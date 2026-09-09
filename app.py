import os
from datetime import datetime, timedelta
from decimal import Decimal
from functools import wraps

from flask import (
    Flask,
    abort,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    session,
    url_for,
    send_file,
)
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from flask_login import LoginManager, current_user, login_required, login_user, logout_user
from sqlalchemy import func, text, inspect
from werkzeug.utils import secure_filename

from config import Config
from models import (
    Cart,
    Category,
    Order,
    OrderItem,
    Product,
    SiteSettings,
    User,
    db
)

app = Flask(__name__)
app.config.from_object(Config)
db.init_app(app)

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"
login_manager.login_message = "กรุณาเข้าสู่ระบบก่อนใช้งาน"
login_manager.login_message_category = "warning"

os.makedirs(app.config["UPLOAD_FOLDER"], exist_ok=True)
os.makedirs(app.config["SLIP_FOLDER"], exist_ok=True)


@login_manager.unauthorized_handler
def unauthorized():
    if request.path.startswith("/admin") or request.path.startswith("/backoffice"):
        return redirect(url_for("backoffice_login", next=request.url))
    return redirect(url_for("login", next=request.url))


@login_manager.user_loader
def load_user(user_id):
    return db.session.get(User, int(user_id))


def admin_required(f):
    @wraps(f)
    @login_required
    def decorated(*args, **kwargs):
        if not current_user.is_admin:
            abort(403)
        return f(*args, **kwargs)

    return decorated


def allowed_file(filename):
    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower() in app.config["ALLOWED_EXTENSIONS"]
    )


def save_upload(file, folder):
    if file and file.filename and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        saved_name = f"{timestamp}_{filename}"
        file.save(os.path.join(folder, saved_name))
        return saved_name
    return None


def expire_pending_orders(order=None):
    now = datetime.now()
    query = Order.query.filter_by(status="pending")
    if order:
        query = query.filter(Order.id == order.id)
    for pending in query.all():
        if pending.payment_expires_at and now > pending.payment_expires_at:
            pending.status = "cancelled"
    db.session.commit()


def get_cart():
    return Cart.query.filter_by(user_id=1).all()


def save_cart(cart):
    pass

# ============================
# POS (ระบบหน้าร้าน) - SINGLE SOURCE ONLY
# ============================

POS_CART = []
PAYMENT_UID = "C3110714"

def build_cart_response():
    cart_items = []
    total = 0.0
    for item in POS_CART:
        product = db.session.get(Product, item["product_id"])
        if not product:
            continue
        cart_items.append({
            "id": product.id,
            "name": product.name,
            "price": float(product.price),
            "stock": product.stock,
            "quantity": item["quantity"],
            "image": url_for("static", filename=f"uploads/{product.image}") if product.image else ""
        })
        total += float(product.price) * item["quantity"]
    return {"cart": cart_items, "total": total}


@app.route("/admin/pos", endpoint="admin_pos")
@admin_required
def admin_pos():
    products = Product.query.filter_by(is_active=True).all()
    categories = Category.query.order_by(Category.id.asc()).all()

    products_data = [{
        "id": p.id,
        "name": p.name,
        "price": float(p.price),
        "stock": p.stock,
        "category_id": p.category_id,
        "category_name": p.category.name if p.category else "ไม่มีหมวดหมู่",
        "image": url_for(
            "static",
            filename=f"uploads/{p.image}"
        ) if p.image else ""
    } for p in products]

    categories_data = [{
        "id": c.id,
        "name": c.name
    } for c in categories]

    return render_template(
        "admin/pos.html",
        products=products_data,
        categories=categories_data
    )


@app.route("/api/admin/cart")
def api_admin_cart():
    return jsonify(build_cart_response())


@app.route("/api/admin/cart/add", methods=["POST"])
def api_admin_cart_add():
    global POS_CART
    data = request.get_json()
    pid = int(data["product_id"])
    qty = int(data.get("quantity", 1))
    product = Product.query.get_or_404(pid)
    if qty > product.stock:
        return jsonify({"success": False, "message": "สินค้าไม่เพียงพอ"})
    for item in POS_CART:
        if item["product_id"] == pid:
            new_qty = item["quantity"] + qty
            if new_qty > product.stock:
                return jsonify({"success": False, "message": "สินค้าไม่เพียงพอ"})
            item["quantity"] = new_qty
            return jsonify({"success": True})
    POS_CART.append({"product_id": pid, "quantity": qty})
    return jsonify({"success": True})


@app.route("/api/admin/cart/update", methods=["POST"])
def api_admin_cart_update():
    global POS_CART
    data = request.get_json()
    pid = int(data["product_id"])
    qty = int(data["quantity"])
    for item in POS_CART:
        if item["product_id"] == pid:
            product = Product.query.get(pid)
            if qty > product.stock:
                qty = product.stock
            item["quantity"] = max(1, qty)
            break
    return jsonify({"success": True})


@app.route("/api/admin/cart/remove", methods=["POST"])
def api_admin_cart_remove():
    global POS_CART
    data = request.get_json()
    pid = int(data["product_id"])
    POS_CART = [i for i in POS_CART if i["product_id"] != pid]
    return jsonify({"success": True})


@app.route("/api/admin/rfid_scan", methods=["POST"])
def admin_rfid_scan():
    global POS_CART
    data = request.get_json()
    uid = data.get("uid", "").replace(" ", "").upper()
    product = Product.query.filter_by(rfid_uid=uid, is_active=True).first()
    if not product:
        return jsonify({"success": False, "message": "ไม่พบสินค้า"})
    for item in POS_CART:
        if item["product_id"] == product.id:
            item["quantity"] += 1
            return jsonify({"success": True, **build_cart_response()})
    POS_CART.append({"product_id": product.id, "quantity": 1})
    return jsonify({"success": True, **build_cart_response()})


@app.route("/api/admin/payment", methods=["POST"])
def api_admin_payment():
    global POS_CART
    if not POS_CART:
        return jsonify({"success": False, "message": "ไม่มีสินค้าในตะกร้า"})
    data = request.get_json()
    admin = User.query.filter_by(role="admin").first()
    order = Order(
        user_id=admin.id,
        status="paid",
        total=0,
        shipping_name=data.get("customer_name", ""),
        shipping_phone=data.get("customer_phone", ""),
        shipping_address=data.get("customer_address", "")
    )
    db.session.add(order)
    db.session.flush()
    total = 0
    for item in POS_CART:
        product = db.session.get(Product, item["product_id"])
        if not product:
            continue
        total += product.price * item["quantity"]
        db.session.add(OrderItem(
            order_id=order.id,
            product_id=product.id,
            quantity=item["quantity"],
            price=product.price
        ))
        product.stock -= item["quantity"]
    order.total = total
    db.session.commit()
    POS_CART.clear()
    return jsonify({"success": True, "message": "ชำระเงินสำเร็จ"})

def cart_count():
    if not current_user.is_authenticated:
        return 0

    return (
        db.session.query(func.sum(Cart.quantity))
        .filter_by(user_id=current_user.id)
        .scalar()
        or 0
    )

def cart_items():

    if not current_user.is_authenticated:
        return [], Decimal("0")

    cart = Cart.query.filter_by(
        user_id=current_user.id
    ).all()

    items = []
    total = Decimal("0")

    for item in cart:

        product = item.product

        if product and product.is_active:

            subtotal = Decimal(str(product.price)) * item.quantity

            items.append({
                "product": product,
                "quantity": item.quantity,
                "subtotal": subtotal
            })

            total += subtotal

    return items, total


@app.context_processor
def inject_globals():

    cart_count = 0

    if current_user.is_authenticated:
        cart_count = Cart.query.filter_by(
            user_id=current_user.id
        ).count()

    return dict(
        cart_count=cart_count,
        categories=Category.query.order_by(Category.name).all()
    )


@app.template_filter("product_image")
def product_image(image):
    upload_path = os.path.join(app.config["UPLOAD_FOLDER"], image)
    if os.path.exists(upload_path):
        return url_for("static", filename=f"uploads/{image}")
    return url_for("static", filename=f"images/{image}")


@app.route("/")
def index():

    category = request.args.get("category", type=int)
    keyword = request.args.get("q", "").strip()

    products = Product.query.filter_by(is_active=True)

    if category:
        products = products.filter_by(category_id=category)

    if keyword:
        products = products.filter(
            Product.name.ilike(f"%{keyword}%")
        )

    products = products.all()

    categories = Category.query.order_by(Category.name).all()

    return render_template(
        "index.html",
        products=products,
        categories=categories,
        selected_category=category,
        keyword=keyword
    )


@app.route("/product/<int:product_id>")
def product_detail(product_id):
    product = db.session.get(Product, product_id)
    if not product or not product.is_active:
        abort(404)
    return render_template("product_detail.html", product=product)


@app.route("/cart")
def cart():
    items, total = cart_items()
    return render_template("cart.html", items=items, total=total)


@app.route("/cart/add/<int:product_id>", methods=["POST"])
@login_required
def add_to_cart(product_id):

    product = Product.query.get_or_404(product_id)

    item = Cart.query.filter_by(
        user_id=current_user.id,
        product_id=product.id
    ).first()

    if item:

        if item.quantity < product.stock:
            item.quantity += 1

    else:

        db.session.add(
            Cart(
                user_id=current_user.id,
                product_id=product.id,
                quantity=1
            )
        )

    db.session.commit()

    flash("เพิ่มสินค้าแล้ว", "success")

    return redirect(url_for("cart"))


@app.route("/cart/update/<int:product_id>", methods=["POST"])
@login_required
def update_cart(product_id):

    qty = int(request.form.get("quantity", 1))

    item = Cart.query.filter_by(
        user_id=current_user.id,
        product_id=product_id
    ).first()

    if item:

        if qty <= 0:
            db.session.delete(item)

        else:
            item.quantity = qty

        db.session.commit()

    return redirect(url_for("cart"))


@app.route("/cart/remove/<int:product_id>", methods=["POST"])
@login_required
def remove_from_cart(product_id):

    item = Cart.query.filter_by(
        user_id=current_user.id,
        product_id=product_id
    ).first()

    if item:
        db.session.delete(item)
        db.session.commit()

    flash("ลบสินค้าแล้ว", "success")

    return redirect(url_for("cart"))


@app.route("/register", methods=["GET", "POST"])
def register():
    if current_user.is_authenticated:
        return redirect(url_for("index"))

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm_password", "")

        if not username or not email or not password:
            flash("กรุณากรอกข้อมูลให้ครบ", "error")
        elif password != confirm:
            flash("รหัสผ่านไม่ตรงกัน", "error")
        elif User.query.filter_by(username=username).first():
            flash("ชื่อผู้ใช้นี้มีอยู่แล้ว", "error")
        elif User.query.filter_by(email=email).first():
            flash("อีเมลนี้มีอยู่แล้ว", "error")
        else:
            user = User(username=username, email=email, role="user")
            user.set_password(password)
            db.session.add(user)
            db.session.commit()
            flash("สมัครสมาชิกสำเร็จ กรุณาเข้าสู่ระบบ", "success")
            return redirect(url_for("login"))

    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("index"))

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = User.query.filter_by(username=username).first()

        if user and user.check_password(password):
            if user.is_admin:
                flash("กรุณาเข้าสู่ระบบผ่านหน้า Backoffice", "warning")
                return render_template("login.html")
            login_user(user)
            next_page = request.args.get("next")
            flash(f"ยินดีต้อนรับ {user.username}", "success")
            return redirect(next_page or url_for("index"))

        flash("ชื่อผู้ใช้หรือรหัสผ่านไม่ถูกต้อง", "error")

    return render_template("login.html")


@app.route("/backoffice/login", methods=["GET", "POST"])
def backoffice_login():
    if current_user.is_authenticated and current_user.is_admin:
        return redirect(url_for("admin_dashboard"))

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = User.query.filter_by(username=username, role="admin").first()

        if user and user.check_password(password):
            login_user(user)
            flash(f"ยินดีต้อนรับ Admin {user.username}", "success")
            return redirect(request.args.get("next") or url_for("admin_dashboard"))

        flash("ชื่อผู้ใช้หรือรหัสผ่านแอดมินไม่ถูกต้อง", "error")

    return render_template("admin/login.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    flash("ออกจากระบบแล้ว", "success")
    return redirect(url_for("index"))


@app.route("/backoffice/logout")
@login_required
def backoffice_logout():
    logout_user()
    flash("ออกจากระบบแอดมินแล้ว", "success")
    return redirect(url_for("backoffice_login"))


@app.route("/test")
def test():
    return "OK"


@app.route("/api/rfid_scan", methods=["POST"])
def api_rfid_scan():

    data = request.get_json()

    if not data:
        return jsonify({"success": False, "message": "No Data"}), 400

    uid = data.get("uid", "").replace(" ", "").upper()

    if uid == "":
        return jsonify({"success": False, "message": "UID Empty"}), 400

    product = Product.query.filter_by(
        rfid_uid=uid,
        is_active=True
    ).first()

    if product is None:
        return jsonify({
            "success": False,
            "message": "ไม่พบสินค้า"
        }), 404

    # ทดลองใช้ user_id = 1 ก่อน
    user_id = 1

    item = Cart.query.filter_by(
        user_id=user_id,
        product_id=product.id
    ).first()

    if item:

        if item.quantity < product.stock:
            item.quantity += 1

    else:

        item = Cart(
            user_id=user_id,
            product_id=product.id,
            quantity=1
        )

        db.session.add(item)

    db.session.commit()

    return jsonify({
        "success": True,
        "product": product.to_dict(),
        "quantity": item.quantity
    })



@app.route("/checkout", methods=["GET", "POST"])
@login_required
def checkout():

    items = []
    total = 0

    cart_items_db = Cart.query.filter_by( user_id=current_user.id).all()

    for c in cart_items_db:
        subtotal = c.product.price * c.quantity
        total += subtotal

        items.append({
            "product": c.product,
            "quantity": c.quantity,
            "subtotal": subtotal
        })

    if not items:
        flash("ตะกร้าว่างเปล่า", "warning")
        return redirect(url_for("index"))

    if request.method == "POST":

        name = request.form.get("shipping_name", "").strip()
        phone = request.form.get("shipping_phone", "").strip()
        address = request.form.get("shipping_address", "").strip()

        if not name or not phone or not address:
            flash("กรุณากรอกข้อมูลจัดส่งให้ครบ", "error")
            return render_template(
                "checkout.html",
                items=items,
                total=total
            )

        order = Order(
            user_id=current_user.id,
            total=total,
            status="pending",
            shipping_name=name,
            shipping_phone=phone,
            shipping_address=address,
            payment_expires_at=datetime.now()
            + timedelta(minutes=app.config["PAYMENT_TIMEOUT_MINUTES"])
        )

        db.session.add(order)
        db.session.flush()

        for item in items:

            product = item["product"]

            if item["quantity"] > product.stock:
                db.session.rollback()
                flash(f"{product.name} มีสินค้าไม่เพียงพอ", "error")
                return redirect(url_for("cart"))

            db.session.add(
                OrderItem(
                    order_id=order.id,
                    product_id=product.id,
                    quantity=item["quantity"],
                    price=product.price
                )
            )

            product.stock -= item["quantity"]

        db.session.commit()

        # ล้างตะกร้า
        Cart.query.filter_by(user_id=current_user.id).delete()
        db.session.commit()

        flash("สร้างคำสั่งซื้อสำเร็จ", "success")
        return redirect(url_for("order_detail", order_id=order.id))

    return render_template(
        "checkout.html",
        items=items,
        total=total
    )


@app.route("/orders/<int:order_id>")
@login_required
def order_detail(order_id):

    order = Order.query.get_or_404(order_id)

    if order.user_id != current_user.id:
        abort(403)

    settings = SiteSettings.get()

    return render_template(
        "order_detail.html",
        order=order,
        settings=settings
    )

@app.route("/orders")
@login_required
def orders():
    expire_pending_orders()

    user_orders = (
        Order.query
        .filter_by(user_id=current_user.id)
        .order_by(Order.created_at.desc())
        .all()
    )

    return render_template(
        "orders.html",
        orders=user_orders
    )


@app.route("/orders/<int:order_id>/upload-slip", methods=["POST"])
@login_required
def upload_slip(order_id):
    order = db.session.get(Order, order_id)
    if not order or order.user_id != current_user.id:
        abort(404)

    expire_pending_orders(order)
    order = db.session.get(Order, order_id)

    if order.status != "pending":
        flash("ไม่สามารถอัปโหลดสลิปสำหรับคำสั่งซื้อนี้ได้", "error")
        return redirect(url_for("order_detail", order_id=order_id))

    if order.is_payment_expired:
        flash("หมดเวลาชำระเงินแล้ว", "error")
        return redirect(url_for("order_detail", order_id=order_id))

    if "payment_slip" not in request.files:
        flash("กรุณาเลือกไฟล์สลิป", "error")
        return redirect(url_for("order_detail", order_id=order_id))

    slip_name = save_upload(request.files["payment_slip"], app.config["SLIP_FOLDER"])
    if not slip_name:
        flash("ไฟล์สลิปไม่ถูกต้อง (รองรับ jpg, png, gif, webp)", "error")
        return redirect(url_for("order_detail", order_id=order_id))

    order.payment_slip = slip_name
    order.status = "awaiting_verification"
    db.session.commit()
    flash("อัปโหลดสลิปสำเร็จ รอแอดมินตรวจสอบ", "success")
    return redirect(url_for("order_detail", order_id=order_id))


# ─── Admin Routes ───────────────────────────────────────────────

@app.route("/admin")
@admin_required
def admin_dashboard():
    today = datetime.now().date()
    month_start = today.replace(day=1)

    today_sales = (
        db.session.query(func.coalesce(func.sum(Order.total), 0))
        .filter(func.date(Order.created_at) == today, Order.status != "cancelled")
        .scalar()
    )
    month_sales = (
        db.session.query(func.coalesce(func.sum(Order.total), 0))
        .filter(func.date(Order.created_at) >= month_start, Order.status != "cancelled")
        .scalar()
    )
    total_orders = Order.query.filter(Order.status != "cancelled").count()
    total_products = Product.query.count()
    pending_orders = Order.query.filter(
        Order.status.in_(["paid", "awaiting_verification"])
    ).count()

    recent_orders = Order.query.order_by(Order.created_at.desc()).limit(5).all()

    return render_template(
        "admin/dashboard.html",
        today_sales=today_sales,
        month_sales=month_sales,
        total_orders=total_orders,
        total_products=total_products,
        pending_orders=pending_orders,
        recent_orders=recent_orders,
    )

@app.route("/admin/products")
@admin_required
def admin_products():
    products = Product.query.order_by(Product.created_at.desc()).all()
    return render_template("admin/products.html", products=products)

@app.route("/admin/categories")
@admin_required
def admin_categories():

    categories = Category.query.order_by(Category.id.asc()).all()

    return render_template(
        "admin/categories.html",
        categories=categories
    )



@app.route("/admin/categories/add", methods=["GET", "POST"])
@admin_required
def admin_category_add():

    if request.method == "POST":

        name = request.form.get("name", "").strip()

        if not name:
            flash("กรุณากรอกชื่อหมวดหมู่", "error")
            return render_template(
                "admin/category_form.html",
                category=None
            )

        exists = Category.query.filter_by(name=name).first()

        if exists:
            flash("ชื่อหมวดหมู่นี้มีอยู่แล้ว", "warning")
            return render_template(
                "admin/category_form.html",
                category=None
            )

        category = Category(
            name=name
        )

        db.session.add(category)
        db.session.commit()

        flash("เพิ่มหมวดหมู่สำเร็จ", "success")

        return redirect(url_for("admin_categories"))

    return render_template(
        "admin/category_form.html",
        category=None
    )


@app.route("/admin/categories/edit/<int:category_id>", methods=["GET", "POST"])
@admin_required
def admin_category_edit(category_id):

    category = db.session.get(Category, category_id)

    if not category:
        abort(404)

    if request.method == "POST":

        name = request.form.get("name", "").strip()

        if not name:
            flash("กรุณากรอกชื่อหมวดหมู่", "error")
            return redirect(
                url_for(
                    "admin_category_edit",
                    category_id=category.id
                )
            )

        exists = Category.query.filter(
            Category.name == name,
            Category.id != category.id
        ).first()

        if exists:
            flash("ชื่อหมวดหมู่นี้มีอยู่แล้ว", "warning")
            return redirect(
                url_for(
                    "admin_category_edit",
                    category_id=category.id
                )
            )

        # บันทึกข้อมูล
        category.name = name

        db.session.commit()

        flash("แก้ไขหมวดหมู่สำเร็จ", "success")

        return redirect(url_for("admin_categories"))

    return render_template(
        "admin/category_form.html",
        category=category
    )



@app.route("/admin/categories/delete/<int:category_id>", methods=["POST"])
@admin_required
def admin_category_delete(category_id):

    category = db.session.get(Category, category_id)

    if not category:
        abort(404)

    if category.products:

        flash(
            "ลบไม่ได้ เพราะยังมีสินค้าอยู่ในหมวดหมู่นี้",
            "warning"
        )

        return redirect(url_for("admin_categories"))

    db.session.delete(category)

    db.session.commit()

    flash("ลบหมวดหมู่สำเร็จ", "success")

    return redirect(url_for("admin_categories"))


@app.route("/admin/products/add", methods=["GET", "POST"])
@admin_required
def admin_product_add():

    categories = Category.query.order_by(
        Category.name.asc()
    ).all()

    if request.method == "POST":

        # =========================
        # ข้อมูลพื้นฐาน
        # =========================

        name = request.form.get(
            "name", ""
        ).strip()

        description = request.form.get(
            "description", ""
        ).strip()

        # =========================
        # RFID
        # =========================

        rfid_uid = request.form.get(
            "rfid_uid", ""
        ).strip().upper()

        if not rfid_uid or rfid_uid == "NONE":
            rfid_uid = None

        else:

            # ตรวจ RFID ซ้ำ
            existing_product = Product.query.filter_by(
                rfid_uid=rfid_uid
            ).first()

            if existing_product:

                flash(
                    f"RFID UID {rfid_uid} "
                    f"ถูกใช้งานกับสินค้า "
                    f"'{existing_product.name}' แล้ว",
                    "error"
                )

                return render_template(
                    "admin/product_form.html",
                    product=None,
                    categories=categories
                )

        # =========================
        # Category
        # =========================

        category_id = request.form.get(
            "category_id"
        )

        if category_id:

            try:
                category_id = int(category_id)

            except ValueError:
                category_id = None

        else:
            category_id = None

        # =========================
        # ตรวจข้อมูล
        # =========================

        if not name or not description:

            flash(
                "กรุณากรอกข้อมูลให้ครบ",
                "error"
            )

            return render_template(
                "admin/product_form.html",
                product=None,
                categories=categories
            )

        # =========================
        # ราคา + Stock
        # =========================

        try:

            price = Decimal(
                request.form.get(
                    "price", "0"
                )
            )

            stock = int(
                request.form.get(
                    "stock", "0"
                )
            )

        except Exception:

            flash(
                "ราคาและจำนวนสินค้าต้องเป็นตัวเลข",
                "error"
            )

            return render_template(
                "admin/product_form.html",
                product=None,
                categories=categories
            )

        # =========================
        # รูปภาพ
        # =========================

        image_filename = "shirt1.jpg"

        if "image" in request.files:

            file = request.files["image"]

            if (
                file
                and file.filename
                and allowed_file(file.filename)
            ):

                filename = secure_filename(
                    file.filename
                )

                timestamp = datetime.now().strftime(
                    "%Y%m%d%H%M%S"
                )

                image_filename = (
                    f"{timestamp}_{filename}"
                )

                file.save(
                    os.path.join(
                        app.config["UPLOAD_FOLDER"],
                        image_filename
                    )
                )

        # =========================
        # สร้าง Product
        # =========================

        product = Product(
            name=name,
            description=description,
            price=price,
            stock=stock,
            image=image_filename,
            rfid_uid=rfid_uid,
            category_id=category_id,
            is_active=True
        )

        # =========================
        # บันทึก
        # =========================

        try:

            db.session.add(product)

            db.session.commit()

            flash(
                "เพิ่มสินค้าสำเร็จ",
                "success"
            )

            return redirect(
                url_for("admin_products")
            )

        except Exception as e:

            db.session.rollback()

            flash(
                f"เพิ่มสินค้าไม่สำเร็จ: {e}",
                "error"
            )

    return render_template(
        "admin/product_form.html",
        product=None,
        categories=categories
    )


@app.route("/admin/products/<int:product_id>/edit", methods=["GET", "POST"])
@admin_required
def admin_product_edit(product_id):

    product = db.session.get(Product, product_id)

    if not product:
        abort(404)

    categories = Category.query.order_by(Category.id.asc()).all()

    if request.method == "POST":

        # =========================
        # ข้อมูลพื้นฐาน
        # =========================

        product.name = request.form.get("name", "").strip()

        product.description = request.form.get(
            "description", ""
        ).strip()

        # =========================
        # ราคา
        # =========================

        try:
            product.price = Decimal(
                request.form.get("price", "0")
            )
        except Exception:
            flash("ราคาต้องเป็นตัวเลข", "error")

            return render_template(
                "admin/product_form.html",
                product=product,
                categories=categories
            )

        # =========================
        # Stock
        # =========================

        try:
            product.stock = int(
                request.form.get("stock", "0")
            )
        except Exception:
            flash("จำนวนสินค้าต้องเป็นตัวเลข", "error")

            return render_template(
                "admin/product_form.html",
                product=product,
                categories=categories
            )

        # =========================
        # RFID UID
        # =========================

        rfid_uid = request.form.get(
            "rfid_uid", ""
        ).strip().upper()

        # ถ้าว่าง หรือ NONE ให้เป็น NULL จริง ๆ
        if not rfid_uid or rfid_uid == "NONE":
            product.rfid_uid = None

        else:

            # ตรวจสอบว่า UID นี้มีสินค้าอื่นใช้อยู่หรือไม่
            existing_product = Product.query.filter(
                Product.rfid_uid == rfid_uid,
                Product.id != product.id
            ).first()

            if existing_product:

                flash(
                    f"RFID UID {rfid_uid} ถูกใช้งานกับสินค้า "
                    f"'{existing_product.name}' แล้ว",
                    "error"
                )

                return render_template(
                    "admin/product_form.html",
                    product=product,
                    categories=categories
                )

            product.rfid_uid = rfid_uid

        # =========================
        # Category
        # =========================

        category_id = request.form.get("category_id")

        if category_id:
            try:
                product.category_id = int(category_id)
            except ValueError:
                product.category_id = None
        else:
            product.category_id = None

        # =========================
        # สถานะสินค้า
        # =========================

        product.is_active = (
            request.form.get("is_active") == "on"
        )

        # =========================
        # รูปภาพ
        # =========================

        if "image" in request.files:

            file = request.files["image"]

            if (
                file
                and file.filename
                and allowed_file(file.filename)
            ):

                filename = secure_filename(
                    file.filename
                )

                timestamp = datetime.now().strftime(
                    "%Y%m%d%H%M%S"
                )

                image_filename = (
                    f"{timestamp}_{filename}"
                )

                file.save(
                    os.path.join(
                        app.config["UPLOAD_FOLDER"],
                        image_filename
                    )
                )

                product.image = image_filename

        # =========================
        # บันทึกฐานข้อมูล
        # =========================

        try:

            db.session.commit()

            flash(
                "แก้ไขสินค้าสำเร็จ",
                "success"
            )

            return redirect(
                url_for("admin_products")
            )

        except Exception as e:

            db.session.rollback()

            flash(
                f"แก้ไขสินค้าไม่สำเร็จ: {e}",
                "error"
            )

    return render_template(
        "admin/product_form.html",
        product=product,
        categories=categories
    )


@app.route("/admin/products/<int:product_id>/delete", methods=["POST"])
@admin_required
def admin_product_delete(product_id):

    product = db.session.get(Product, product_id)

    if not product:
        abort(404)

    try:

        # ลบรูป (ถ้าไม่ใช่รูป default)
        if product.image and product.image != "shirt1.jpg":

            image_path = os.path.join(
                app.config["UPLOAD_FOLDER"],
                product.image
            )

            if os.path.exists(image_path):
                os.remove(image_path)

        db.session.delete(product)
        db.session.commit()

        flash("ลบสินค้าสำเร็จ", "success")

    except Exception as e:

        db.session.rollback()
        flash(f"ลบสินค้าไม่สำเร็จ : {e}", "error")

    return redirect(url_for("admin_products"))


@app.route("/admin/orders")
@admin_required
def admin_orders():
    status_filter = request.args.get("status", "")
    query = Order.query
    if status_filter:
        query = query.filter_by(status=status_filter)
    orders_list = query.order_by(Order.created_at.desc()).all()
    return render_template("admin/orders.html", orders=orders_list, status_filter=status_filter)


@app.route("/admin/orders/<int:order_id>", methods=["GET", "POST"])
@admin_required
def admin_order_detail(order_id):
    order = db.session.get(Order, order_id)
    if not order:
        abort(404)

    if request.method == "POST":
        action = request.form.get("action")
        if action == "approve_slip" and order.status == "awaiting_verification":
            order.status = "paid"
            for item in order.items:
                item.product.stock -= item.quantity
            db.session.commit()
            flash("อนุมัติสลิปและยืนยันการชำระเงินแล้ว", "success")
        elif action == "reject_slip" and order.status == "awaiting_verification":
            order.status = "pending"
            order.payment_slip = None
            order.payment_expires_at = datetime.now() + timedelta(
                minutes=app.config["PAYMENT_TIMEOUT_MINUTES"]
            )
            db.session.commit()
            flash("ปฏิเสธสลิป ลูกค้าสามารถอัปโหลดใหม่ได้", "warning")
        else:
            new_status = request.form.get("status")
            if new_status in Order.STATUS_LABELS:
                order.status = new_status
                db.session.commit()
                flash("อัปเดตสถานะคำสั่งซื้อแล้ว", "success")
        return redirect(url_for("admin_order_detail", order_id=order_id))

    return render_template("admin/order_detail.html", order=order)


@app.route("/admin/settings", methods=["GET", "POST"])
@admin_required
def admin_settings():
    settings = SiteSettings.get()

    if request.method == "POST":
        settings.promptpay_number = request.form.get("promptpay_number", "").strip()
        settings.promptpay_name = request.form.get("promptpay_name", "").strip()

        if "qr_image" in request.files:
            qr_name = save_upload(request.files["qr_image"], app.config["UPLOAD_FOLDER"])
            if qr_name:
                settings.qr_image = qr_name

        db.session.commit()
        flash("บันทึกการตั้งค่าชำระเงินแล้ว", "success")
        return redirect(url_for("admin_settings"))

    return render_template("admin/settings.html", settings=settings)


@app.route("/admin/reports")
@admin_required
def admin_reports():
    period = request.args.get("period", "daily")
    today = datetime.now().date()

    if period == "monthly":
        start_date = today.replace(day=1) - timedelta(days=365)
        sales_data = (
            db.session.query(
                func.strftime("%Y-%m", Order.created_at).label("period"),
                func.sum(Order.total).label("total"),
                func.count(Order.id).label("count"),
            )
            .filter(Order.created_at >= start_date, Order.status != "cancelled")
            .group_by("period")
            .order_by("period")
            .all()
        )
        title = "รายงานยอดขายรายเดือน"
    else:
        start_date = today - timedelta(days=30)
        sales_data = (
            db.session.query(
                func.date(Order.created_at).label("period"),
                func.sum(Order.total).label("total"),
                func.count(Order.id).label("count"),
            )
            .filter(Order.created_at >= start_date, Order.status != "cancelled")
            .group_by(func.date(Order.created_at))
            .order_by("period")
            .all()
        )
        title = "รายงานยอดขายรายวัน"

    grand_total = sum(float(row.total or 0) for row in sales_data)
    grand_count = sum(row.count for row in sales_data)

    return render_template(
        "admin/reports.html",
        sales_data=sales_data,
        period=period,
        title=title,
        grand_total=grand_total,
        grand_count=grand_count,
    )


@app.route("/admin/reports/export")
@admin_required
def admin_reports_export():

    period = request.args.get("period", "all")
    selected_date = request.args.get("date")
    selected_month = request.args.get("month")

    query = Order.query.filter(Order.status != "cancelled")

    # รายวัน
    if period == "daily" and selected_date:
        try:
            date_obj = datetime.strptime(selected_date, "%Y-%m-%d").date()

            start_datetime = datetime.combine(date_obj, datetime.min.time())
            end_datetime = start_datetime + timedelta(days=1)

            query = query.filter(
                Order.created_at >= start_datetime,
                Order.created_at < end_datetime
            )

        except ValueError:
            pass

    # รายเดือน
    elif period == "monthly" and selected_month:
        try:
            month_obj = datetime.strptime(selected_month, "%Y-%m")

            start_datetime = month_obj.replace(day=1)

            if start_datetime.month == 12:
                end_datetime = start_datetime.replace(
                    year=start_datetime.year + 1,
                    month=1
                )
            else:
                end_datetime = start_datetime.replace(
                    month=start_datetime.month + 1
                )

            query = query.filter(
                Order.created_at >= start_datetime,
                Order.created_at < end_datetime
            )

        except ValueError:
            pass

    orders = query.order_by(Order.created_at.desc()).all()

    wb = Workbook()

    ws = wb.active
    ws.title = "Sales Report"

    headers = [
        "Order ID",
        "วันที่",
        "ลูกค้า",
        "สินค้า",
        "หมวดหมู่",
        "จำนวน",
        "ราคาต่อชิ้น",
        "รวม",
        "สถานะ"
    ]

    for i, h in enumerate(headers, start=1):

        cell = ws.cell(row=1, column=i)
        cell.value = h
        cell.font = Font(bold=True)

    row = 2

    for order in orders:

        for item in order.items:

            ws.cell(row=row, column=1).value = order.id

            ws.cell(row=row, column=2).value = (
                order.created_at.strftime("%d/%m/%Y %H:%M")
            )

            ws.cell(row=row, column=3).value = (
                order.user.username
                if order.user
                else "-"
            )

            ws.cell(row=row, column=4).value = item.product.name

            ws.cell(row=row, column=5).value = (
                item.product.category.name
                if item.product.category
                else "-"
            )

            ws.cell(row=row, column=6).value = item.quantity

            ws.cell(row=row, column=7).value = float(item.price)

            ws.cell(row=row, column=8).value = (
                float(item.price) * item.quantity
            )

            ws.cell(row=row, column=9).value = order.status

            row += 1

    # ปรับความกว้างคอลัมน์
    for column_cells in ws.columns:

        length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        ws.column_dimensions[
            column_cells[0].column_letter
        ].width = length + 5

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    # ตั้งชื่อไฟล์ตามประเภท
    if period == "daily" and selected_date:

        filename = f"Sales_Report_{selected_date}.xlsx"

    elif period == "monthly" and selected_month:

        filename = f"Sales_Report_{selected_month}.xlsx"

    else:

        filename = "Sales_Report_All.xlsx"

    return send_file(
        output,
        download_name=filename,
        as_attachment=True,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def migrate_db():

    inspector = inspect(db.engine)

    with db.engine.begin() as conn:

        # ---------- orders ----------
        if "orders" in inspector.get_table_names():

            cols = {c["name"] for c in inspector.get_columns("orders")}

            if "payment_slip" not in cols:
                conn.execute(text(
                    "ALTER TABLE orders ADD COLUMN payment_slip VARCHAR(255)"
                ))

            if "payment_expires_at" not in cols:
                conn.execute(text(
                    "ALTER TABLE orders ADD COLUMN payment_expires_at DATETIME"
                ))

            # ---------- categories ----------
            if "categories" not in inspector.get_table_names():
                conn.execute(text("""
                    CREATE TABLE categories (
                        id INTEGER PRIMARY KEY AUTO_INCREMENT,
                        name VARCHAR(100) UNIQUE NOT NULL
                )
            """))

        # ---------- products ----------
        if "products" in inspector.get_table_names():

            cols = {c["name"] for c in inspector.get_columns("products")}

            if "rfid_uid" not in cols:
                conn.execute(text(
                    "ALTER TABLE products ADD COLUMN rfid_uid VARCHAR(20)"
                ))

            if "category_id" not in cols:
                conn.execute(text(
                    "ALTER TABLE products ADD COLUMN category_id INTEGER"
                ))

def init_db():
    with app.app_context():
        db.create_all()
        migrate_db()

        admin = User.query.filter_by(username="admin").first()
        if not admin:
            admin = User(
                username="admin",
                email="admin@shop.com",
                role="admin"
            )
            admin.set_password("123456")
            db.session.add(admin)
        else:
            admin.set_password("123456")

        db.session.commit()

        SiteSettings.get()

        # ==========================
        # เพิ่มหมวดหมู่เริ่มต้น
        # ==========================
        if Category.query.count() == 0:
            db.session.add_all([
                Category(name="Arduino"),
                Category(name="ESP32"),
                Category(name="Sensor"),
                Category(name="Relay"),
                Category(name="Motor"),
                Category(name="Display"),
                Category(name="Module"),
                Category(name="Power"),
                Category(name="IC"),
                Category(name="Passive"),
                Category(name="Cable"),
                Category(name="Tool")
            ])
            db.session.commit()

        # ==========================
        # เพิ่มสินค้าเริ่มต้น
        # ==========================
        if Product.query.count() == 0:
            sample_products = [
                Product(
                    name="เสื้อยืดคอกลม สีขาว",
                    description="เสื้อยืดผ้าฝ้าย 100% สวมใส่สบาย เหมาะสำหรับทุกโอกาส",
                    price=Decimal("299.00"),
                    stock=50,
                    image="shirt1.jpg",
                ),
                Product(
                    name="เสื้อโปโล สีน้ำเงิน",
                    description="เสื้อโปโลคุณภาพดี ดีไซน์เรียบหรู ใส่ได้ทั้งทำงานและเที่ยว",
                    price=Decimal("450.00"),
                    stock=30,
                    image="shirt2.jpg",
                ),
                Product(
                    name="เสื้อเชิ้ตลายตาราง",
                    description="เสื้อเชิ้ตแขนยาวลายตาราง สไตล์สมาร์ทแคชชวล",
                    price=Decimal("590.00"),
                    stock=20,
                    image="shirt3.jpg",
                ),
            ]

            db.session.add_all(sample_products)
            db.session.commit()


if __name__ == "__main__":
    with app.app_context():
        print("DATABASE URI:", app.config["SQLALCHEMY_DATABASE_URI"])
        print("DATABASE PATH:", db.engine.url.database)

    if os.environ.get("WERKZEUG_RUN_MAIN") == "true":
        init_db()

    app.run(host="0.0.0.0", debug=True, port=5000)
